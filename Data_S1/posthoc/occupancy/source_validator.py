"""Independent source reconstruction and manual CR1 check; no analyze.py import.

Read-only original CSV/XLSX inputs. Output only independent_validation.json.
Uses csv/openpyxl for extraction, numpy.linalg.lstsq and a manually assembled
gene-cluster sandwich, and scipy.stats.t for intervals. No statsmodels/patsy.
"""
from pathlib import Path
import csv
import hashlib
import json
import math
import re
import sys
from datetime import datetime, timezone
import numpy as np
from scipy.stats import t
from openpyxl import load_workbook

PACKAGE_ROOT = next(p for p in Path(__file__).resolve().parents if p.name == 'Data_S1')
sys.path.insert(0, str(PACKAGE_ROOT))
from input_support import expected_count, input_path

HERE = Path(__file__).resolve().parent
ROOT = Path(__file__).resolve().parents[2]
RES = ROOT / 'analysis/results'
PATHS = {
    'properties': RES / 'step0_properties.csv',
    'binding': RES / 'step1_binding.csv',
    'expression': RES / 'step1_expression.csv',
    'rpf': RES / 'step2_rpf_counts.csv',
    'catalogue': ROOT / 'analysis/data/Mudge2022_SupplementaryTables.xlsx',
}


def sha(path):
    h = hashlib.sha256()
    with path.open('rb') as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def csv_index(path, key):
    out = {}
    with path.open(newline='') as handle:
        for row in csv.DictReader(line for line in handle if not line.startswith('#')):
            assert row[key] not in out, (path.name, row[key])
            out[row[key]] = row
    return out


def stable(value):
    # Preserve a suffix such as _PAR_Y while removing the numeric version.
    return re.sub(r'\.\d+', '', value)


def num(value):
    try:
        return float(value)
    except (ValueError, TypeError):
        return float('nan')


def reconstruct():
    properties = csv_index(PATHS['properties'], 'orf_name')
    binding = csv_index(PATHS['binding'], 'orf_name')
    expression = csv_index(PATHS['expression'], 'orf_name')
    rpf = csv_index(PATHS['rpf'], 'orf_name')
    hosts = csv_index(PATHS['host_metadata'], 'enst')
    workbook = load_workbook(PATHS['catalogue'], read_only=True, data_only=True)
    iterator = workbook['S2. PHASE I Ribo-seq ORFs'].iter_rows(values_only=True)
    headers = next(iterator)
    i_orf, i_gene = headers.index('orf_name'), headers.index('gene_id')
    genes = {}
    for row in iterator:
        if row[i_orf] is None:
            continue
        assert row[i_orf] not in genes
        assert row[i_gene] is not None
        genes[row[i_orf]] = stable(str(row[i_gene]))
    workbook.close()
    population = {name: row for name, row in properties.items()
                  if row['source_sheet'] == 'S2' and num(row['aa_length']) <= 100}
    assert len(population) == expected_count('main_orfs')
    analyses = {}
    for cell in ['K562', 'HepG2']:
        selected = []
        for name in sorted(population):
            row = population[name]
            enst = stable(row['transcript'])
            ensg = genes[name]
            assert ensg == hosts[enst]['ensg']
            values = [num(expression[name][f'tpm_{cell}_polyA']),
                      num(binding[name][f'{cell}_nrbp_tx']),
                      num(rpf[name][f'{cell}_rpf_in']),
                      num(row['aa_length']), num(row['transcript_gc']),
                      num(hosts[enst]['spliced_length_nt'])]
            tpm, bound, footprints, aa, gc, length = values
            if not (tpm >= 1 and all(math.isfinite(v) for v in values)):
                continue
            assert aa > 0 and length > 0 and bound >= 0 and footprints >= 0
            selected.append((name, enst, ensg, tpm, bound, footprints, aa, gc, length))
        numeric = np.asarray([row[3:] for row in selected], dtype=float)
        tpm, bound, footprints, aa, gc, length = numeric.T
        X = np.column_stack([np.ones(len(selected)), np.log1p(bound),
                             np.log(aa), gc, np.log1p(tpm)])
        analyses[cell] = {
            'names': [row[0] for row in selected],
            'tx': [row[1] for row in selected],
            'genes': [row[2] for row in selected],
            'X': X, 'y': np.log2(footprints + 1), 'log_length': np.log(length),
        }
    return analyses


def manual_fit(X, y, groups):
    n, p = X.shape
    assert np.linalg.matrix_rank(X) == p
    coefficients = np.linalg.lstsq(X, y, rcond=None)[0]
    residuals = y - X @ coefficients
    _, R = np.linalg.qr(X, mode='reduced')
    inverse_R = np.linalg.inv(R)
    bread = inverse_R @ inverse_R.T
    scores = {}
    for label, xi, error in zip(groups, X, residuals):
        if label not in scores:
            scores[label] = np.zeros(p)
        scores[label] += xi * error
    G = len(scores)
    meat = sum((np.outer(score, score) for score in scores.values()), np.zeros((p, p)))
    correction = (G / (G - 1)) * ((n - 1) / (n - p))
    covariance = correction * bread @ meat @ bread
    se = float(np.sqrt(covariance[1, 1]))
    beta = float(coefficients[1])
    critical = float(t.ppf(0.975, G-1))
    reduced = np.delete(X, 1, axis=1)
    reduced_residual = y - reduced @ np.linalg.lstsq(reduced, y, rcond=None)[0]
    sst = float(((y-y.mean())**2).sum())
    sse = float(residuals @ residuals)
    increment = float((reduced_residual @ reduced_residual-sse)/sst)
    return {
        'n_ORFs': n, 'n_genes': G, 'p_design_columns': p,
        'beta': beta, 'SE': se, 'ci95_lo': beta-critical*se,
        'ci95_hi': beta+critical*se, 'p': float(2*t.sf(abs(beta/se), G-1)),
        'incremental_R2_RBP': increment, 'full_R2': 1-sse/sst,
        'df_resid_inference': G-1, 'CR1_correction': correction,
        'condition_number_X': float(np.linalg.cond(X)),
        'all_coefficients': coefficients.tolist(),
    }


def main():
    PATHS['host_metadata'] = input_path('host_transcript_metadata')
    source_hashes = {key: {'path': str(path), 'sha256': sha(path)} for key, path in PATHS.items()}
    data = reconstruct()
    # Only final comparison reads the principal implementation's output.
    reference = json.loads((HERE/'results.json').read_text())
    records = []
    suppressor_identities = []
    for cell, a in data.items():
        fits = {}
        row_hash = hashlib.sha256('\n'.join(a['names']).encode()).hexdigest()
        diagnostic = next(v for v in reference['diagnostics'] if v['cell'] == cell)
        assert row_hash == diagnostic['ORF_ids_sha256'], 'Different ORF row set'
        for model, X in [('original', a['X']),
                         ('plus_log_length', np.column_stack([a['X'], a['log_length']]))]:
            result = manual_fit(X, a['y'], a['genes'])
            fits[model] = result
            original = next(v for v in reference['model_results']
                            if v['cell'] == cell and v['model'] == model and v['cluster'] == 'gene')
            differences = {}
            for key in ['beta', 'SE', 'ci95_lo', 'ci95_hi', 'p',
                        'incremental_R2_RBP', 'full_R2']:
                differences[key] = result[key]-original[key]
                assert np.isclose(result[key], original[key], atol=1e-10, rtol=1e-9), (cell, model, key)
            assert result['n_ORFs'] == original['n_ORFs']
            assert result['n_genes'] == original['n_genes']
            result.update(cell=cell, model=model, cluster='gene',
                          n_transcripts=len(set(a['tx'])), ORF_ids_sha256=row_hash,
                          differences_vs_main=differences)
            records.append(result)
        length_on_baseline = np.linalg.lstsq(a['X'], a['log_length'], rcond=None)[0]
        slope_length = fits['plus_log_length']['all_coefficients'][-1]
        delta_observed = fits['original']['beta']-fits['plus_log_length']['beta']
        delta_identity = float(length_on_baseline[1]*slope_length)
        assert np.isclose(delta_observed, delta_identity, atol=1e-10)
        suppressor_identities.append({
            'cell': cell,
            'length_slope_in_full_model': slope_length,
            'RBP_slope_predicting_log_length_adjusted_for_original_covariates': float(length_on_baseline[1]),
            'baseline_beta_minus_length_adjusted_beta': delta_observed,
            'omitted_variable_identity_product': delta_identity,
            'interpretation': 'Algebraic adjustment identity, not a causal decomposition',
        })
    for key, path in PATHS.items():
        assert sha(path) == source_hashes[key]['sha256']
    out = {
        'completed_utc': datetime.now(timezone.utc).isoformat(),
        'method': 'Independent csv/openpyxl reconstruction, numpy least squares, manually assembled gene-cluster CR1 covariance with G-1 t intervals. No analyze.py, statsmodels or patsy import.',
        'sources': source_hashes,
        'validation_script_sha256': sha(Path(__file__)),
        'status': 'PASS', 'checked_models': records,
        'coefficient_increase_algebra': suppressor_identities,
    }
    (HERE/'independent_validation.json').write_text(json.dumps(out, indent=2, allow_nan=False)+'\n')
    for r in records:
        print(r['cell'], r['model'], 'n=', r['n_ORFs'], 'beta=', r['beta'],
              'SE=', r['SE'], 'CI=', [r['ci95_lo'], r['ci95_hi']],
              'incremental_R2=', r['incremental_R2_RBP'])
    print('PASS: independent point estimates, CR1 SE, t intervals, p, R2 and exact ORF sets')
    print(json.dumps(suppressor_identities, indent=2))


if __name__ == '__main__':
    main()

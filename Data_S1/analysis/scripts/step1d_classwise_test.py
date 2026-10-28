"""(ii) Functional-class RBP joint model — prespecified one-shot.

Functional classes are fixed from Van Nostrand 2020 (Nature, MOESM3) BEFORE
looking at results — this is the literature-based class assignment, not a
post-hoc grouping.

Pre-specified directional hypotheses (per cell line):
  Translation regulation        -> pLDDT coefficient positive (foldability)
  P-body / stress granules      -> catGRANULE coefficient positive (LLPS prone)
  Splicing regulation           -> null (no effect on microprotein properties)
  3' end processing             -> null
(RNA modification listed as neutral; not in the 4-test set.)

Per ORF: count of unique class-RBPs binding the host transcript exons
(strand-matched, 1bp), per cell line, per class.

Joint model (coding-host MAIN pop, per cell, per outcome, OLS+HC3):
  outcome ~ n_translation + n_PB_SG + n_splicing + n_3end
            + log(AA_len) + GC + log1p(TPM)

Multiple comparison: the two positive hypotheses each have two cell-specific
tests (four effect+ rows in total), using Bonferroni α = 0.05 / 4 = 0.0125.
A separate Bonferroni correction covers all 16 reported class coefficients.
Positive = prior-direction effect significant IN BOTH CELL LINES.

Runs in .venv_ranges. Writes results/step1d_classwise_results.csv.
"""
from pathlib import Path
from collections import defaultdict

import numpy as np
import pandas as pd
import pyranges as pr
import openpyxl
import statsmodels.formula.api as smf
from statsmodels.stats.multitest import multipletests

ROOT = Path(__file__).resolve().parent.parent
STEP0 = ROOT / "results" / "step0_properties.csv"
REGION = ROOT / "results" / "step1c_region_binding.csv"  # has host_coding & merged metadata
EXPR = ROOT / "results" / "step1_expression.csv"
EXONS = ROOT / "results" / "v35_orf_exons.parquet"
VN_XLSX = ROOT / "data" / "VanNostrand2020_RBP_functions.xlsx"
ECLIP_DIR = ROOT / "data" / "eclip"
OUT = ROOT / "results" / "step1d_classwise_results.csv"
CELLS = ["K562", "HepG2"]

# prespecified classes (column names in MOESM3 Sheet 1, row 1)
CLASSES = {
    "translation":   ("Translation regulation",   "pLDDT",      "effect+"),
    "pb_sg":         ("P-body / stress granules", "catGRANULE", "effect+"),
    "splicing":      ("Splicing regulation",      "any",        "null"),
    "three_end":     ("3' end processing",        "any",        "null"),
}


def read_step0():
    nc = 0
    for line in open(STEP0):
        if line.startswith("#"):
            nc += 1
        else:
            break
    return pd.read_csv(STEP0, skiprows=nc)


def load_vn_classes():
    """RBP name -> set of class keys (using our 4 pre-specified classes)."""
    wb = openpyxl.load_workbook(VN_XLSX, read_only=True, data_only=True)
    ws = wb["Sheet 1"]
    rows = list(ws.iter_rows(values_only=True))
    sub = rows[1]
    # column index per Van Nostrand category name
    col = {sub[i]: i for i in range(len(sub)) if sub[i]}
    members = {k: set() for k in CLASSES}
    for r in rows[2:]:
        if not r[0]:
            continue
        rbp = str(r[0]).strip().upper()
        for key, (catname, _, _) in CLASSES.items():
            if catname not in col:
                raise KeyError(f"Van Nostrand category {catname!r} not found in MOESM3")
            v = r[col[catname]]
            if v in (1, "1", True):
                members[key].add(rbp)
    return members


def load_peaks_with_rbp(cell):
    rows = []
    for f in sorted((ECLIP_DIR / cell).glob("*.bed.gz")):
        rbp = f.name.split("__")[0]
        d = pd.read_csv(f, sep="\t", header=None, compression="gzip",
                        usecols=[0, 1, 2, 5],
                        names=["Chromosome", "Start", "End", "Strand"],
                        dtype={0: str, 1: int, 2: int, 5: str})
        d["rbp"] = rbp.upper()
        rows.append(d)
    return pr.PyRanges(pd.concat(rows, ignore_index=True))


def main():
    s0 = read_step0()
    s0["enst"] = s0["transcript"].str.split(".").str[0]
    expr = pd.read_csv(EXPR)[["orf_name", "tpm_K562_polyA", "tpm_HepG2_polyA"]]
    region = pd.read_csv(REGION)[["orf_name", "host_coding"]]
    df = s0.merge(expr, on="orf_name").merge(region, on="orf_name")
    df["l_aalen"] = np.log(df["aa_length"])
    df["gc"] = df["transcript_gc"]

    # --- Van Nostrand classes ---
    members = load_vn_classes()
    print("Van Nostrand 2020 class membership (full RBP set, n=356):")
    for k, (name, _, _) in CLASSES.items():
        print(f"  {k:13s} ({name}): {len(members[k])} RBPs")

    # --- exon map (versionless ENST) ---
    exon_df = pd.read_parquet(EXONS)
    exon_pr = pr.PyRanges(exon_df)

    # --- intersect per cell, per RBP -> per-ENST set of binding RBPs ---
    eclip_rbps = {}  # cell -> set of RBPs actually in our eCLIP data
    enst_rbps = {}   # cell -> {enst: set(rbps)}
    for cell in CELLS:
        print(f"\nIntersecting {cell} peaks per RBP ...", flush=True)
        peaks = load_peaks_with_rbp(cell)
        eclip_rbps[cell] = set(peaks.df["rbp"].unique())
        j = exon_pr.join(peaks, strandedness="same").df
        e2r = j.groupby("enst")["rbp"].agg(set).to_dict()
        enst_rbps[cell] = e2r
        print(f"  {cell}: {len(eclip_rbps[cell])} unique RBPs; "
              f"{len(e2r)} transcripts with >=1 binding", flush=True)

    # --- class coverage in our eCLIP data ---
    print("\n=== prespecified class coverage in our eCLIP set ===")
    print(f"{'class':14s} {'VN_total':>9s} {'K562_eclip':>11s} {'HepG2_eclip':>12s} "
          f"{'both_eclip':>11s}")
    for k, (name, _, _) in CLASSES.items():
        vn = members[k]
        c_k = vn & eclip_rbps["K562"]
        c_h = vn & eclip_rbps["HepG2"]
        c_b = c_k & c_h
        print(f"{k:14s} {len(vn):>9d} {len(c_k):>11d} {len(c_h):>12d} {len(c_b):>11d}")
    # also overall coverage: how many of our eCLIP RBPs are in Van Nostrand at all
    vn_all = set().union(*members.values())  # union of our 4 classes only
    print("\noverall: of our eCLIP RBPs, how many are listed in MOESM3 at all?")
    moesm_all = set()
    wb = openpyxl.load_workbook(VN_XLSX, read_only=True, data_only=True)
    for r in list(wb["Sheet 1"].iter_rows(values_only=True))[2:]:
        if r[0]:
            moesm_all.add(str(r[0]).strip().upper())
    for cell in CELLS:
        print(f"  {cell}: {len(eclip_rbps[cell] & moesm_all)} / {len(eclip_rbps[cell])} "
              f"({len(eclip_rbps[cell] - moesm_all)} not in MOESM3: "
              f"{sorted(eclip_rbps[cell] - moesm_all)[:8]}...)")

    # --- per-ORF class binding counts ---
    for cell in CELLS:
        for k in CLASSES:
            cls_rbps = members[k] & eclip_rbps[cell]
            df[f"{cell}_n_{k}"] = df["enst"].map(
                lambda e: len(enst_rbps[cell].get(e, set()) & cls_rbps)
            )

    # --- joint model: coding-host MAIN ---
    base = df[(df.source_sheet == "S2") & (df.aa_length <= 100) & (df.host_coding)].copy()
    rows = []
    OUTCOMES = {"mean_plddt": "pLDDT", "catgranule2_llps_score": "catGRANULE"}
    PRIOR_FOR = {
        ("translation", "pLDDT"): "effect+",
        ("translation", "catGRANULE"): "neutral",
        ("pb_sg", "pLDDT"): "neutral",
        ("pb_sg", "catGRANULE"): "effect+",
        ("splicing", "pLDDT"): "null",
        ("splicing", "catGRANULE"): "null",
        ("three_end", "pLDDT"): "null",
        ("three_end", "catGRANULE"): "null",
    }

    for cell in CELLS:
        d = base.copy()
        d["l1p_tpm"] = np.log1p(d[f"tpm_{cell}_polyA"])
        cls_cols = {k: f"{cell}_n_{k}" for k in CLASSES}
        need = list(cls_cols.values()) + ["l1p_tpm", "l_aalen", "gc",
                                          "mean_plddt", "catgranule2_llps_score"]
        n_before = len(d)
        d = d.dropna(subset=need)
        # stable names for formula
        for k in CLASSES:
            d[f"n_{k}"] = d[cls_cols[k]]
        print(f"\n{cell}: analysis n={len(d)} (coding-host main {n_before}; "
              f"dropped NaN TPM: {n_before-len(d)})")
        rhs = " + ".join(f"n_{k}" for k in CLASSES) + " + l_aalen + gc + l1p_tpm"
        for outcome, olabel in OUTCOMES.items():
            m = smf.ols(f"{outcome} ~ {rhs}", data=d).fit(cov_type="HC3")
            for k in CLASSES:
                term = f"n_{k}"
                beta = m.params[term]
                lo, hi = m.conf_int().loc[term]
                rows.append(dict(cell=cell, outcome=olabel, vn_class=k,
                                 prior=PRIOR_FOR[(k, olabel)],
                                 beta=beta, ci_lo=lo, ci_hi=hi,
                                 p_raw=m.pvalues[term], R2=m.rsquared, n=len(d)))

    res = pd.DataFrame(rows)
    # Bonferroni over the four cell-specific tests where prior is effect+
    pre_reg = res[res["prior"] == "effect+"].copy()
    pre_reg_idx = pre_reg.index.tolist()
    if len(pre_reg_idx):
        res.loc[pre_reg_idx, "p_bonf_pre"] = multipletests(
            res.loc[pre_reg_idx, "p_raw"], method="bonferroni")[1]
    # also overall Bonferroni over all 16 coefficients reported
    res["p_bonf_all"] = multipletests(res["p_raw"], method="bonferroni")[1]
    res["p_BH"] = multipletests(res["p_raw"], method="fdr_bh")[1]

    OUT.parent.mkdir(parents=True, exist_ok=True)
    res.to_csv(OUT, index=False)
    print("\n================ CLASS-WISE JOINT MODEL — result table ================")
    pd.set_option("display.width", 200, "display.max_columns", 25)
    print(res.round({"beta": 3, "ci_lo": 3, "ci_hi": 3,
                     "p_raw": 4, "p_BH": 4, "p_bonf_pre": 4, "p_bonf_all": 4,
                     "R2": 3}).to_string(index=False))

    # final positive-judgment per prespecified hypothesis (both-cell replicate)
    print("\n=== prespecified hypothesis judgment (Bonferroni alpha=0.05/4=0.0125, "
          "BOTH cells same-sign and significant) ===")
    pre_reg2 = res[res["prior"] == "effect+"]   # re-filter after adding p_bonf_pre
    for (cls, oc), g in pre_reg2.groupby(["vn_class", "outcome"]):
        if g["cell"].nunique() < 2:
            continue
        beta_signs = np.sign(g["beta"])
        same_sign = beta_signs.nunique() == 1 and beta_signs.iloc[0] > 0  # expecting +
        both_sig = (g["p_raw"] < 0.05).all()
        bonf_sig = (g["p_bonf_pre"] < 0.05).all()
        verdict = "POSITIVE" if (same_sign and both_sig and bonf_sig) else "null"
        print(f"  {cls:13s} -> {oc:10s}  verdict: {verdict}  "
              f"(beta K562={g[g.cell=='K562'].beta.iloc[0]:+.4f}, "
              f"HepG2={g[g.cell=='HepG2'].beta.iloc[0]:+.4f})")
    print(f"\nWrote {OUT}")


if __name__ == "__main__":
    main()

